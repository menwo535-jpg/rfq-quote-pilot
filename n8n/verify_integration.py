"""Execute synthetic cases in a real, isolated n8n CLI instance and local RFQ API.

Install requirements-test.txt plus n8n 2.38.7 separately, then pass its bin/n8n.
Only this script's new temporary instance is used. Existing n8n data is untouched.
"""
from __future__ import annotations

import argparse
import base64
import hashlib
import io
import json
import os
import secrets
import socket
import subprocess
import sys
import tempfile
import threading
import time
from datetime import datetime, timezone
from pathlib import Path

from build_workflow import ROOT, workflow
from local_api import create_server
from make_request import make_request
from openpyxl import load_workbook

MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
DRAFT = "Draft awaiting human approval"
REVIEW = "Unresolved lines - review required"


def text_payload(rows, catalogue):
    return {
        "rfq": {"format": "txt", "content": "ITEMS\n" + rows + "\nEND ITEMS"},
        "catalogue_csv": "code,description,unit,unit_price,currency,price_version\n" + catalogue,
        "include_xlsx": True,
    }


def cases():
    return [
        ("sample_pdf", json.loads(make_request(ROOT / "samples/rfq.pdf", ROOT / "samples/catalogue.csv", include_xlsx=True)), REVIEW, {"USD": "32.50"}, 4),
        ("matched_currencies", text_payload("001|Cable|2.5|M\n002|插头|3|EA", "001,Cable,M,1.234,USD,v1\n002,插头,EA,2.5,EUR,v1\n"), DRAFT, {"USD": "3.09", "EUR": "7.50"}, 0),
        ("zero_price", text_payload("0001|Sample|2|EA", "0001,Sample,EA,0,USD,v1\n"), DRAFT, {"USD": "0.00"}, 0),
        ("duplicate_lines", text_payload("A|Cable|1|M\nA|Cable|1|M", "A,Cable,M,5,USD,v1\n"), REVIEW, {}, 2),
        ("unknown_code", text_payload("X|Unknown|1|EA", "A,Cable,M,5,USD,v1\n"), REVIEW, {}, 1),
        ("invalid_catalogue", dict(text_payload("A|Cable|1|M", "A,Cable,M,5,USD,v1\n"), catalogue_csv="bad,headers\n1,2\n"), "ERROR", None, None),
        ("xlsx_precision_limit", text_payload("A|Bulk|999999999|EA", "A,Bulk,EA,999999999,USD,v1\n"), "ERROR", None, None),
    ]


def execution_from_log(text):
    # CLI can emit migrations/status text even with --rawOutput. Parse the full
    # execution object, not a success message or the CLI's exit code alone.
    decoder = json.JSONDecoder()
    for offset, character in enumerate(text):
        if character != "{" or (offset and text[offset - 1] != "\n"):
            continue
        try:
            value, _ = decoder.raw_decode(text[offset:])
        except ValueError:
            continue
        if isinstance(value, dict) and isinstance(value.get("data"), dict) and "resultData" in value["data"]:
            return value
    raise AssertionError("n8n did not return a parseable execution object; inspect local raw logs")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--node", default="node")
    parser.add_argument("--n8n-bin", type=Path, required=True)
    parser.add_argument("--out", type=Path, required=True)
    parser.add_argument("--case", help="run one case while diagnosing a failure")
    args = parser.parse_args()
    args.out.mkdir(parents=True, exist_ok=True)
    selected = [case for case in cases() if not args.case or case[0] == args.case]
    if not selected:
        parser.error("unknown case")
    token = secrets.token_hex(24)
    server = create_server(token, 0)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    summary = {"started_utc": datetime.now(timezone.utc).isoformat(), "synthetic_only": True, "cases": []}
    try:
        # TemporaryDirectory uses native Python deletion for its own exact path.
        # No user's existing instance or credentials are read/exported/changed.
        with tempfile.TemporaryDirectory(prefix="rfq-n8n-verify-") as temporary:
            private = Path(temporary)
            instance = private / "instance"
            instance.mkdir()
            with socket.socket() as sock:
                sock.bind(("127.0.0.1", 0))
                broker_port = sock.getsockname()[1]
            # Do not inherit another n8n instance's DB/storage/hooks settings.
            instance_prefixes = ("N8N_", "DB_", "EXECUTIONS_", "QUEUE_", "BINARY_DATA_", "EXTERNAL_HOOK_FILES")
            env = {key: value for key, value in os.environ.items() if not key.startswith(instance_prefixes)}
            env.update({
                "N8N_USER_FOLDER": str(instance), "DB_TYPE": "sqlite",
                "N8N_ENCRYPTION_KEY": secrets.token_hex(32),
                "N8N_DIAGNOSTICS_ENABLED": "false", "N8N_VERSION_NOTIFICATIONS_ENABLED": "false",
                "N8N_TEMPLATES_ENABLED": "false", "N8N_PERSONALIZATION_ENABLED": "false",
                "N8N_LISTEN_ADDRESS": "127.0.0.1", "N8N_HOST": "127.0.0.1",
                "N8N_RUNNERS_BROKER_LISTEN_ADDRESS": "127.0.0.1",
                "N8N_RUNNERS_BROKER_PORT": str(broker_port),
                "N8N_DEFAULT_BINARY_DATA_MODE": "default", "N8N_LOG_LEVEL": "info",
            })

            def cli(label, arguments):
                start = time.monotonic()
                process = subprocess.run([args.node, str(args.n8n_bin.resolve()), *arguments],
                                         cwd=ROOT, env=env, capture_output=True, text=True,
                                         encoding="utf-8", errors="replace", timeout=180)
                combined = process.stdout + "\n" + process.stderr
                # The token is a new local test secret, never a user credential.
                # Refuse to persist it even if an upstream CLI logs request data.
                combined = combined.replace(token, "[LOCAL_TEST_TOKEN_REDACTED]")
                (args.out / (label + ".log")).write_text(combined, encoding="utf-8")
                return process.returncode, combined, round(time.monotonic() - start, 3)

            code, version, _ = cli("version", ["--version"])
            if code or version.strip() != "2.38.7":
                raise AssertionError("This check requires n8n 2.38.7; inspect version.log")
            summary["n8n_version"] = version.strip()
            summary["node_version"] = subprocess.check_output([args.node, "--version"], text=True).strip()
            credential_id = "rfqSyntheticHeader01"
            credentials = private / "credential.json"
            credentials.write_text(json.dumps([{
                "id": credential_id, "name": "RFQ local API", "type": "httpHeaderAuth",
                "data": {"name": "Authorization", "value": "Bearer " + token},
            }]), encoding="utf-8")
            code, output, _ = cli("import-credential", ["import:credentials", "--input=" + str(credentials)])
            if code or "Successfully imported" not in output:
                raise AssertionError("Local test credential import failed; inspect local log")
            credentials.unlink()

            imported = []
            for name, payload, *_ in selected:
                document = workflow(payload, url=f"http://127.0.0.1:{server.server_port}/quote", credential_id=credential_id)
                document["id"] = "rfq" + name.replace("_", "")
                imported.append(document)
            workflow_file = private / "workflows.json"
            workflow_file.write_text(json.dumps(imported, ensure_ascii=False), encoding="utf-8")
            code, output, _ = cli("import-workflows", ["import:workflow", "--input=" + str(workflow_file)])
            if code or "Successfully imported" not in output:
                raise AssertionError("Workflow import failed; inspect local log")

            for document, (name, payload, expected, subtotals, review_count) in zip(imported, selected):
                code, raw, elapsed = cli(name, ["execute", "--id=" + document["id"], "--rawOutput"])
                execution = execution_from_log(raw)
                result = execution["data"]["resultData"]
                run_data = result.get("runData", {})
                error = result.get("error")
                actual_branches = [branch for branch in (DRAFT, REVIEW) if branch in run_data]
                detail = {"case": name, "seconds": elapsed, "cli_exit_code": code}
                if expected == "ERROR":
                    assert error and not actual_branches, (name, "error must stop before review/approval outputs")
                    assert result.get("lastNodeExecuted") == "Prepare draft", (name, "wrong failing node")
                    assert str(error.get("httpCode")) == "422", (name, "expected actual HTTP validation failure", error)
                    detail.update({"outcome": "validation_error", "http_status": 422, "file_outputs": 0})
                else:
                    assert not error and actual_branches == [expected], (name, error, actual_branches)
                    response = run_data["Prepare draft"][0]["data"]["main"][0][0]["json"]
                    quote = response["quote"]
                    assert quote["review_count"] == review_count, (name, quote)
                    assert quote["matched_subtotals_by_currency"] == subtotals, (name, quote)
                    assert quote["complete_totals_by_currency"] == (subtotals if expected == DRAFT else None)
                    endpoint_bytes = base64.b64decode(response["artifacts"]["xlsx_base64"], validate=True)
                    item = run_data[expected][0]["data"]["main"][0][0]
                    binary = item["binary"]["workbook"]
                    assert binary["mimeType"] == MIME and binary["fileName"] == "quote-draft.xlsx"
                    decoded = base64.b64decode(binary["data"], validate=True)
                    assert decoded == endpoint_bytes, (name, "n8n changed or double-encoded the workbook")
                    book = load_workbook(io.BytesIO(decoded), data_only=False)
                    assert book.sheetnames == ["Quote", "Sources"]
                    for sheet in book:
                        assert all(cell.data_type != "f" for row in sheet for cell in row)
                    values = list(book["Quote"].values)
                    if name == "matched_currencies":
                        assert any("001" in row for row in values)
                        assert any("插头" in row for row in values)
                    if name == "zero_price":
                        assert any("0001" in row for row in values)
                    book.close()
                    (args.out / (name + ".xlsx")).write_bytes(decoded)
                    detail.update({"outcome": quote["status"], "branch": expected, "review_count": review_count,
                                   "matched_subtotals": subtotals, "xlsx_bytes": len(decoded),
                                   "xlsx_sha256": hashlib.sha256(decoded).hexdigest()})
                detail["passed"] = True
                summary["cases"].append(detail)
                print(json.dumps(detail, ensure_ascii=True), flush=True)
                (args.out / "validation.json").write_text(json.dumps(summary, indent=2) + "\n", encoding="utf-8")
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=5)
        summary["finished_utc"] = datetime.now(timezone.utc).isoformat()
        summary["all_passed"] = len(summary["cases"]) == len(selected)
        summary["local_api_stopped"] = not thread.is_alive()
        (args.out / "validation.json").write_text(json.dumps(summary, indent=2) + "\n", encoding="utf-8")


if __name__ == "__main__":
    main()
