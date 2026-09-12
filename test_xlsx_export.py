"""Inspect actual saved XLSX files with an independent reader and real HTTP."""
import base64
import io
import json
import subprocess
import sys
import tempfile
import threading
import unittest
import zipfile
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from local_api import create_server
from make_request import make_request
from quote import parse_pages, prepare_quote, run
import test_local_api
from xlsx_export import render_xlsx

ROOT = Path(__file__).parent


def make_quote(lines, products):
    return prepare_quote(parse_pages(['ITEMS\n' + '\n'.join(lines) + '\nEND ITEMS']), products)


def product(code='A', description='Cable', unit='M', price='1.23', currency='USD'):
    return dict(code=code, description=description, unit=unit, unit_price=price,
                currency=currency, price_version='v1')


def read_workbook(data):
    return load_workbook(io.BytesIO(data), data_only=False)


class XlsxExportTests(unittest.TestCase):
    def test_real_pdf_generates_review_workbook_and_source_digests(self):
        with tempfile.TemporaryDirectory() as folder:
            output = Path(folder)
            result = run(ROOT / 'samples/rfq.pdf', ROOT / 'samples/catalogue.csv', output, include_xlsx=True)
            workbook = load_workbook(output / 'quote-draft.xlsx')
            self.assertEqual(workbook.sheetnames, ['Quote', 'Sources'])
            sheet = workbook['Quote']
            self.assertEqual(sheet['B5'].value, 'REVIEW_REQUIRED')
            self.assertEqual((sheet['B6'].value, sheet['B7'].value), (6, 4))
            self.assertEqual(sheet['E6'].value, 'USD')
            self.assertEqual(sheet['F6'].value, 32.5)
            self.assertIsNone(sheet['G6'].value)
            self.assertEqual(sheet.auto_filter.ref, 'A11:K17')
            self.assertEqual(sheet.freeze_panes, 'C12')
            for index, expected in enumerate(result['rows'], 12):
                self.assertEqual(sheet.cell(index, 2).value, expected['code'])
                self.assertEqual(sheet.cell(index, 6).value, expected['status'])
                for col, key in ((7, 'reason'), (9, 'currency'), (11, 'price_version')):
                    if expected[key] == '':
                        self.assertIsNone(sheet.cell(index, col).value)
                if expected['status'] == 'REVIEW':
                    self.assertIsNone(sheet.cell(index, 8).value)
                    self.assertIsNone(sheet.cell(index, 10).value)
                    self.assertEqual(sheet.cell(index, 7).value, expected['reason'])
            self.assertEqual(workbook['Sources']['B5'].value, 'rfq.pdf')
            self.assertEqual(workbook['Sources']['C5'].value, result['sources']['rfq']['sha256'])
            self.assertEqual(workbook['Sources']['B6'].value, 'catalogue.csv')
            self.assertEqual(workbook['Sources']['C6'].value, result['sources']['approved_catalogue']['sha256'])

    def test_rounding_zero_price_and_currencies_remain_separate(self):
        quote = make_quote(['A|Cable|1|M', 'B|Part|4|EA', 'C|Cap|2|EA'], [
            product(price='1.005'), product('B', 'Part', 'EA', '0', 'EUR'),
            product('C', 'Cap', 'EA', '2.345', 'CNY')])
        sheet = read_workbook(render_xlsx(quote))['Quote']
        self.assertEqual([sheet[f'J{r}'].value for r in range(12, 15)], [1.01, 0, 4.69])
        self.assertEqual([sheet[f'E{r}'].value for r in range(6, 9)], ['CNY', 'EUR', 'USD'])
        self.assertEqual([sheet[f'F{r}'].value for r in range(6, 9)], [4.69, 0, 1.01])
        self.assertEqual([sheet[f'G{r}'].value for r in range(6, 9)], [4.69, 0, 1.01])
        self.assertEqual(sheet['H13'].data_type, 'n')
        self.assertEqual(sheet['B5'].value, 'DRAFT_FOR_APPROVAL')

    def test_leading_zero_identifiers_and_four_decimal_inputs(self):
        quote = make_quote(['000123|电缆|2.1234|M'], [product('000123', '电缆', price='3.4567')])
        sheet = read_workbook(render_xlsx(quote))['Quote']
        self.assertEqual((sheet['B12'].value, sheet['B12'].data_type), ('000123', 's'))
        self.assertEqual(sheet['C12'].value, '电缆')
        self.assertEqual(Decimal(str(sheet['D12'].value)), Decimal('2.1234'))
        self.assertEqual(Decimal(str(sheet['H12'].value)), Decimal('3.4567'))
        self.assertEqual(Decimal(str(sheet['J12'].value)), Decimal(quote['rows'][0]['line_total']))

    def test_formula_like_and_url_text_are_literal_cells(self):
        quote = make_quote(['=1+1|https://example.com/规格😀|1|M'], [product()])
        raw = render_xlsx(quote)
        workbook = read_workbook(raw)
        self.assertEqual(workbook['Quote']['B12'].value, '=1+1')
        self.assertEqual(workbook['Quote']['C12'].value, 'https://example.com/规格😀')
        for sheet in workbook:
            for row in sheet:
                for cell in row:
                    self.assertNotIn(cell.data_type, ('f', 'e'))
                    self.assertIsNone(cell.hyperlink)
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            self.assertFalse(any('externalLinks' in name or 'vbaProject' in name for name in archive.namelist()))

    def test_invalid_quantity_is_preserved_as_text_without_price(self):
        quote = make_quote(['A|Cable|999999999999999999999999999999|M'], [product()])
        sheet = read_workbook(render_xlsx(quote))['Quote']
        self.assertEqual(sheet['D12'].value, '999999999999999999999999999999')
        self.assertEqual(sheet['D12'].data_type, 's')
        self.assertIsNone(sheet['H12'].value)
        self.assertIsNone(sheet['J12'].value)
        self.assertEqual(sheet['F6'].value, 'No priced lines')
        self.assertIsNone(sheet['G6'].value)

    def test_one_review_withholds_every_currency_complete_total(self):
        quote = make_quote(['A|Cable|1|M', 'B|Part|1|EA', 'Z|Unknown|1|EA'],
                           [product(), product('B', 'Part', 'EA', '3', 'EUR')])
        sheet = read_workbook(render_xlsx(quote))['Quote']
        self.assertEqual(sheet['B7'].value, 1)
        self.assertEqual([sheet['F6'].value, sheet['F7'].value], [3, 1.23])
        self.assertIsNone(sheet['G6'].value)
        self.assertIsNone(sheet['G7'].value)

    def test_large_line_amount_is_rejected_without_affecting_core_quote(self):
        quote = make_quote(['A|Cable|999999999|M'], [product(price='999999999')])
        self.assertEqual(quote['complete_totals_by_currency'], {'USD': '999999998000000001.00'})
        with self.assertRaisesRegex(ValueError, 'precision limit'):
            render_xlsx(quote)

    def test_large_sum_is_checked_even_when_each_line_fits(self):
        quote = make_quote(['A|Cable|100000000|M', 'B|Cable|100000000|M'],
                           [product(price='60000'), product('B', price='60000')])
        with self.assertRaisesRegex(ValueError, 'USD subtotal exceeds XLSX precision'):
            render_xlsx(quote)

    def test_cell_length_limit_rejects_instead_of_silently_truncating(self):
        long_text = 'x' * 32767
        quote = make_quote(['A|' + long_text + '|1|M'], [product()])
        self.assertEqual(read_workbook(render_xlsx(quote))['Quote']['C12'].value, long_text)
        quote['rows'][0]['description'] += 'x'
        with self.assertRaisesRegex(ValueError, '32767-character'):
            render_xlsx(quote)

    def test_maximum_item_count_is_not_truncated(self):
        quote = make_quote([f'{i:06}|Cable|1|M' for i in range(500)],
                           [product(f'{i:06}') for i in range(500)])
        sheet = read_workbook(render_xlsx(quote))['Quote']
        self.assertEqual(sheet['B6'].value, 500)
        self.assertEqual(sheet['B511'].value, '000499')
        self.assertEqual(sheet['F6'].value, 615)
        self.assertEqual(sheet.auto_filter.ref, 'A11:K511')

    def test_cli_flag_and_review_exit_code(self):
        with tempfile.TemporaryDirectory() as folder:
            output = Path(folder) / 'result'
            completed = subprocess.run([sys.executable, str(ROOT / 'quote.py'), str(ROOT / 'samples/rfq.pdf'),
                                        str(ROOT / 'samples/catalogue.csv'), '--out', str(output), '--xlsx'],
                                       capture_output=True, text=True)
            self.assertEqual(completed.returncode, 2, completed.stderr)
            self.assertEqual(load_workbook(output / 'quote-draft.xlsx')['Quote']['B7'].value, 4)

    def test_legacy_cli_has_no_xlsx_and_xlsx_failure_does_not_touch_outputs(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            run(ROOT / 'samples/rfq.txt', ROOT / 'samples/catalogue.csv', root / 'legacy')
            self.assertFalse((root / 'legacy/quote-draft.xlsx').exists())
            (root / 'huge.txt').write_text('ITEMS\nA|Cable|999999999|M\nEND ITEMS')
            (root / 'catalogue.csv').write_text('code,description,unit,unit_price,currency,price_version\nA,Cable,M,999999999,USD,v1\n')
            with self.assertRaisesRegex(ValueError, 'precision limit'):
                run(root / 'huge.txt', root / 'catalogue.csv', root / 'must-not-exist', include_xlsx=True)
            self.assertFalse((root / 'must-not-exist').exists())


class XlsxHTTPTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.server = create_server(test_local_api.TEST_TOKEN, 0)
        cls.thread = threading.Thread(target=cls.server.serve_forever, daemon=True)
        cls.thread.start()

    @classmethod
    def tearDownClass(cls):
        cls.server.shutdown()
        cls.server.server_close()
        cls.thread.join(timeout=5)

    def setUp(self):
        # Reuse the existing real HTTP client, not a mocked build_response call.
        self.client = test_local_api.LocalAPITests()
        self.client.server = self.server

    def test_real_http_pdf_returns_readable_xlsx_matching_quote(self):
        payload = self.client.payload(pdf=True)
        payload['include_xlsx'] = True
        status, response, _ = self.client.request(payload)
        self.assertEqual(status, 200)
        raw = base64.b64decode(response['artifacts']['xlsx_base64'], validate=True)
        sheet = read_workbook(raw)['Quote']
        self.assertEqual(sheet['B7'].value, response['quote']['review_count'])
        self.assertEqual(sheet['F6'].value, 32.5)
        self.assertIsNone(sheet['G6'].value)
        self.assertEqual(response['quote']['status'], 'REVIEW_REQUIRED')

    def test_default_and_false_keep_original_http_shape(self):
        payload = self.client.payload()
        old = self.client.request(payload)[1]
        payload['include_xlsx'] = False
        explicit_false = self.client.request(payload)[1]
        self.assertEqual(old, explicit_false)
        self.assertEqual(set(old['artifacts']), {'csv_utf8', 'review_html'})

    def test_flag_is_boolean_and_precision_error_is_actionable(self):
        payload = self.client.payload()
        for value in ('true', 1, None, [], {}):
            payload['include_xlsx'] = value
            self.assertEqual(self.client.request(payload)[0], 422)
        payload = {'rfq': {'format': 'txt', 'content': 'ITEMS\nA|Cable|999999999|M\nEND ITEMS'},
                   'catalogue_csv': 'code,description,unit,unit_price,currency,price_version\nA,Cable,M,999999999,USD,v1\n',
                   'include_xlsx': True}
        status, response, _ = self.client.request(payload)
        self.assertEqual(status, 422)
        self.assertIn('precision limit', response['error'])
        payload['include_xlsx'] = False
        self.assertEqual(self.client.request(payload)[0], 200)

    def test_request_builder_opts_in_without_changing_default(self):
        args = (ROOT / 'samples/rfq.pdf', ROOT / 'samples/catalogue.csv')
        self.assertNotIn('include_xlsx', json.loads(make_request(*args)))
        payload = json.loads(make_request(*args, include_xlsx=True))
        self.assertIs(payload['include_xlsx'], True)
        self.assertEqual(self.client.request(payload)[0], 200)


if __name__ == '__main__':
    unittest.main()
